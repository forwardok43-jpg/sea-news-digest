#!/usr/bin/env python3
from __future__ import annotations

import asyncio
import calendar
import html
import json
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, replace
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import parse_qs, quote, quote_plus, urlparse
from zoneinfo import ZoneInfo

import feedparser
import requests
import trafilatura
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from googlenewsdecoder import gnews_decoder_async


TZ = ZoneInfo("Asia/Shanghai")
WECOM_WEBHOOK_URL = os.getenv("WECOM_WEBHOOK_URL", "").strip()
OPENROUTER_API_KEY = os.getenv("OPENROUTER_API_KEY", "").strip()
OPENROUTER_URL = "https://openrouter.ai/api/v1/chat/completions"
CLOUDFLARE_ACCOUNT_ID = os.getenv("CLOUDFLARE_ACCOUNT_ID", "").strip()
CLOUDFLARE_API_TOKEN = os.getenv("CLOUDFLARE_API_TOKEN", "").strip()
DRY_RUN = os.getenv("DRY_RUN", "0") == "1"
LOG_PREVIEW = os.getenv("LOG_PREVIEW", "0") == "1"
ENABLE_AI_SUMMARY = os.getenv("ENABLE_AI_SUMMARY", "0") == "1"



def env_int(name: str, default: int) -> int:
    try:
        return int(os.getenv(name, str(default)))
    except ValueError:
        return default


FETCH_WORKERS = max(1, env_int("FETCH_WORKERS", 10))
TRANSLATE_WORKERS = max(1, env_int("TRANSLATE_WORKERS", 4))
SUMMARY_BATCH_SIZE = max(1, env_int("SUMMARY_BATCH_SIZE", 8))
SUMMARY_WORKERS = max(1, env_int("SUMMARY_WORKERS", 2))
OVERVIEW_MIN_CHARS = max(100, env_int("OVERVIEW_MIN_CHARS", 300))
OVERVIEW_MAX_CHARS = max(OVERVIEW_MIN_CHARS, env_int("OVERVIEW_MAX_CHARS", 500))
MAX_ITEMS = max(1, env_int("MAX_ITEMS", 200))
MAX_PER_SOURCE = max(1, env_int("MAX_PER_SOURCE", 4))
LOOKBACK_HOURS = max(1, env_int("LOOKBACK_HOURS", 24))
DIGEST_ITEMS = max(3, env_int("DIGEST_ITEMS", 80))
USER_AGENT = "Mozilla/5.0 (compatible; sea-news-digest/1.0; +https://github.com/forwardok43-jpg/sea-news-digest)"
SOCIAL_DOMAINS = ("facebook.com", "instagram.com", "x.com", "twitter.com", "youtube.com", "tiktok.com")

DEFAULT_SUMMARY_MODELS = (
    "thinkingmachines/inkling:free,"
    "dots-studio/dots-3-note-preview:free,"
    "nvidia/nemotron-3-super-120b-a12b:free,"
    "liquid/lfm-2.5-2.6b:free,"
    "qwen/qwen3.8-27b:free,"
    "google/gemma-4-31b-it:free,"
    "openrouter/free"
)
SUMMARY_MODELS = [
    item.strip()
    for item in os.getenv("SUMMARY_MODELS", DEFAULT_SUMMARY_MODELS).split(",")
    if item.strip()
]

COUNTRIES = [
    ("🇮🇩", "印度尼西亚", "Indonesia", "id", "ID"),
    ("🇻🇳", "越南", "Vietnam", "vi", "VN"),
    ("🇹🇭", "泰国", "Thailand", "th", "TH"),
    ("🇵🇭", "菲律宾", "Philippines", "en", "PH"),
    ("🇲🇾", "马来西亚", "Malaysia", "ms", "MY"),
    ("🇸🇬", "新加坡", "Singapore", "en", "SG"),
    ("🇰🇭", "柬埔寨", "Cambodia", "km", "KH"),
    ("🇲🇲", "缅甸", "Myanmar", "my", "MM"),
    ("🇱🇦", "老挝", "Laos", "lo", "LA"),
    ("🇧🇳", "文莱", "Brunei", "en", "BN"),
    ("🇹🇱", "东帝汶", "Timor-Leste", "pt", "TL"),
]

SOURCES: list[dict] = []

for flag, name_zh, search_name, lang, country_code in COUNTRIES:
    SOURCES.append(
        {
            "name": f"{flag} {name_zh}｜政治安全",
            "query": f"{search_name} politics OR election OR parliament OR military OR defense",
            "lang": lang,
            "country": country_code,
            "source_country": name_zh,
        }
    )
    SOURCES.append(
        {
            "name": f"{flag} {name_zh}｜经济产业",
            "query": f"{search_name} economy OR trade OR investment OR infrastructure",
            "lang": lang,
            "country": country_code,
            "source_country": name_zh,
        }
    )
    SOURCES.append(
        {
            "name": f"{flag} {name_zh}｜综合热点",
            "query": search_name,
            "lang": lang,
            "country": country_code,
            "source_country": name_zh,
        }
    )

SOURCES.extend(
    [
        {"name": "🌐 中国—东盟关系", "query": "China ASEAN relations OR summit OR trade OR security", "lang": "en", "country": "SG", "source_country": "中国—东盟", "require_asean": True},
        {"name": "🌐 美国—东盟关系", "query": "United States ASEAN relations OR summit OR defense OR trade", "lang": "en", "country": "SG", "source_country": "美国—东盟", "require_asean": True},
        {"name": "🌐 日本—东盟关系", "query": "Japan ASEAN relations OR summit OR trade OR security", "lang": "en", "country": "SG", "source_country": "日本—东盟", "require_asean": True},
        {"name": "🌐 欧盟—东盟关系", "query": "European Union ASEAN relations OR summit OR trade OR investment", "lang": "en", "country": "SG", "source_country": "欧盟—东盟", "require_asean": True},
        {"name": "🌐 印度—东盟关系", "query": "India ASEAN relations OR summit OR trade OR security", "lang": "en", "country": "SG", "source_country": "印度—东盟", "require_asean": True},
        {"name": "🌐 韩国—东盟关系", "query": "South Korea ASEAN relations OR summit OR trade OR investment", "lang": "en", "country": "SG", "source_country": "韩国—东盟", "require_asean": True},
        {"name": "🌐 俄罗斯—东盟关系", "query": "Russia ASEAN relations OR summit OR energy OR defense", "lang": "en", "country": "SG", "source_country": "俄罗斯—东盟", "require_asean": True},
        {"name": "🌐 澳新—东盟关系", "query": "Australia OR New Zealand ASEAN relations OR summit OR defense", "lang": "en", "country": "SG", "source_country": "澳大利亚/新西兰—东盟", "require_asean": True},
        {"name": "🤝 东盟内部双边关系", "query": "ASEAN bilateral relations OR cooperation OR dispute OR summit", "lang": "en", "country": "SG", "source_country": "东盟内部", "require_asean": True},
        {"name": "🛡️ 东盟安全与南中国海", "query": "ASEAN security OR South China Sea OR military exercise OR defense", "lang": "en", "country": "SG", "source_country": "东盟地区安全", "require_asean": True},
        {"name": "🤝 中国与东南亚经贸安全", "query": "China Southeast Asia trade OR investment OR infrastructure OR security", "lang": "en", "country": "SG", "source_country": "中国—东南亚", "require_asean": True},
        {"name": "🛡️ 美国印太与东南亚", "query": "United States Indo-Pacific Southeast Asia defense OR trade OR alliance", "lang": "en", "country": "SG", "source_country": "美国—东南亚", "require_asean": True},
        {"name": "💼 欧盟与东南亚经贸", "query": "European Union Southeast Asia trade OR investment OR supply chain", "lang": "en", "country": "SG", "source_country": "欧盟—东南亚", "require_asean": True},
        {"name": "🛡️ 日本与东南亚安全经济", "query": "Japan Southeast Asia security OR economy OR investment OR defense", "lang": "en", "country": "SG", "source_country": "日本—东南亚", "require_asean": True},
        {"name": "💼 韩国与东南亚合作", "query": "South Korea Southeast Asia trade OR investment OR cooperation", "lang": "en", "country": "SG", "source_country": "韩国—东南亚", "require_asean": True},
        {"name": "🛡️ 印度与东南亚关系", "query": "India Southeast Asia trade OR security OR Indo-Pacific OR ASEAN", "lang": "en", "country": "SG", "source_country": "印度—东南亚", "require_asean": True},
        {"name": "🛡️ 俄罗斯与东南亚关系", "query": "Russia Southeast Asia energy OR defense OR ASEAN", "lang": "en", "country": "SG", "source_country": "俄罗斯—东南亚", "require_asean": True},
        {"name": "🌐 全球经济与东盟", "query": "global economy OR trade OR supply chain OR investment ASEAN Southeast Asia", "lang": "en", "country": "SG", "source_country": "全球—东盟", "require_asean": True},
        {"name": "🌐 东盟官方", "url": "https://asean.org/feed/", "lang": "en", "country": "SG", "source_country": "东盟官方", "require_asean": True},
        {"name": "🌐 The Diplomat", "url": "https://thediplomat.com/feed/", "lang": "en", "country": "SG", "source_country": "亚太地区", "require_asean": True},
    ]
)

@dataclass(frozen=True)
class Article:
    region: str
    title: str
    link: str
    published: datetime
    publisher: str
    source_country: str = ""
    source_lang: str = ""
    content: str = ""
    summary: str = ""
    title_zh: str = ""


def clean_text(value: object) -> str:
    text = html.unescape(str(value or ""))
    return re.sub(r"\s+", " ", text).strip()


def clean_title(title: str, publisher: str) -> str:
    title = clean_text(title)
    if publisher:
        suffix = f" - {publisher}"
        if title.casefold().endswith(suffix.casefold()):
            title = title[: -len(suffix)].rstrip()
    return title


def truncate_text(text: str, limit: int) -> str:
    text = clean_text(text)
    if len(text) <= limit:
        return text
    return text[: limit - 1].rstrip() + "…"


def chinese_char_count(text: str) -> int:
    return len(re.findall(r"[\u4e00-\u9fff]", text))

def normalize_summary(summary: str, article: Article) -> str:
    text = clean_text(summary)
    text = re.sub(r"^(摘要|总结|简讯)[：:]\s*", "", text)
    text = text.strip("\"'“”")
    text = text.rstrip("。！？!?，,；;：:")
    if chinese_char_count(text) < 4:
        return ""
    if len(text) > 18:
        text = text[:18].rstrip("。！？!?，,；;：:")
    return text


def build_google_news_url(query: str, lang: str, country: str) -> str:
    ceid = f"{country}:{lang}"
    return (
        "https://news.google.com/rss/search"
        f"?q={quote_plus(query)}"
        f"&hl={lang}&gl={country}&ceid={ceid}"
    )


def build_bing_news_url(query: str) -> str:
    return (
        "https://www.bing.com/news/search"
        f"?q={quote_plus(query)}&format=rss&setlang=en-US&cc=US"
    )


def normalize_feed_link(link: str) -> str:
    link = clean_text(link)
    if "bing.com/news/apiclick" not in link:
        return link
    query = parse_qs(urlparse(link).query)
    direct_url = (query.get("url") or [""])[0]
    return direct_url or link

def get_feed(session: requests.Session, url: str):
    last_error = None
    for attempt in range(3):
        try:
            response = session.get(url, timeout=25)
            response.raise_for_status()
            return feedparser.parse(response.content)
        except Exception as exc:
            last_error = exc
            if attempt < 2:
                time.sleep(1.5 * (attempt + 1))
    raise RuntimeError(f"Failed to fetch RSS: {url}") from last_error


def entry_datetime(entry) -> datetime | None:
    for key in ("published_parsed", "updated_parsed"):
        value = entry.get(key)
        if value:
            return datetime.fromtimestamp(calendar.timegm(value), tz=timezone.utc)
    return None


def entry_publisher(entry) -> str:
    source = entry.get("source")
    if isinstance(source, dict):
        return clean_text(source.get("title"))
    return clean_text(getattr(source, "title", ""))


def is_asean_related(title: str) -> bool:
    value = title.casefold()
    terms = (
        "asean",
        "southeast asia",
        "south china sea",
        "indonesia",
        "vietnam",
        "thailand",
        "philippines",
        "malaysia",
        "singapore",
        "cambodia",
        "myanmar",
        "laos",
        "brunei",
        "timor-leste",
        "mekong",
    )
    return any(term in value for term in terms)

def fetch_source(session: requests.Session, source: dict) -> list[Article]:
    url = source.get("url") or build_google_news_url(source["query"], source.get("lang", "en"), source.get("country", "US"))
    try:
        feed = get_feed(session, url)
    except Exception as primary_error:
        if source.get("url") or "news.google.com" not in url:
            raise
        fallback_url = build_bing_news_url(source["query"])
        print(f"[FALLBACK] {source['name']}: Google RSS failed, trying Bing RSS")
        try:
            feed = get_feed(session, fallback_url)
        except Exception:
            raise primary_error

    if getattr(feed, "bozo", False):
        print(f"[WARN] RSS parse issue: {source['name']}")

    cutoff = datetime.now(timezone.utc) - timedelta(hours=LOOKBACK_HOURS)
    articles: list[Article] = []

    for entry in feed.entries:
        published = entry_datetime(entry)
        if published is None or published < cutoff:
            continue

        publisher = entry_publisher(entry)
        title = clean_title(clean_text(entry.get("title")), publisher)
        link = normalize_feed_link(entry.get("link"))
        if not title or not link:
            continue
        if any(domain in publisher.casefold() for domain in SOCIAL_DOMAINS):
            continue
        if source.get("require_asean") and not is_asean_related(title):
            continue

        articles.append(
            Article(
                region=source["name"],
                title=title,
                link=link,
                published=published,
                publisher=publisher,
                source_country=source.get("source_country", source.get("name", "")),
                source_lang=source.get("lang", "en"),
            )
        )

    articles.sort(key=lambda item: item.published, reverse=True)
    return articles[:MAX_PER_SOURCE]


def deduplication_key(article: Article) -> str:
    return re.sub(r"[\W_]+", "", article.title.casefold())


def collect_articles(session: requests.Session):
    articles: list[Article] = []
    failures: list[str] = []

    def worker(source: dict):
        local_session = requests.Session()
        local_session.headers.update({"User-Agent": USER_AGENT})
        return source, fetch_source(local_session, source)

    with ThreadPoolExecutor(max_workers=min(FETCH_WORKERS, len(SOURCES))) as executor:
        futures = {executor.submit(worker, source): source for source in SOURCES}
        for future in as_completed(futures):
            source = futures[future]
            try:
                _, items = future.result()
                print(f"[OK] {source['name']}: {len(items)}")
                articles.extend(items)
            except Exception as exc:
                message = f"{source['name']}: {exc}"
                print(f"[FAIL] {message}")
                failures.append(message)

    articles.sort(key=lambda item: item.published, reverse=True)
    unique_articles: list[Article] = []
    seen: set[str] = set()

    for article in articles:
        key = deduplication_key(article)
        if not key or key in seen:
            continue
        seen.add(key)
        unique_articles.append(article)

    return unique_articles[:MAX_ITEMS], failures

def select_digest_articles(articles: list[Article], limit: int) -> list[Article]:
    grouped: dict[str, list[Article]] = {source["name"]: [] for source in SOURCES}

    for article in sorted(articles, key=lambda item: item.published, reverse=True):
        grouped.setdefault(article.region, []).append(article)

    selected: list[Article] = []
    while len(selected) < limit:
        made_progress = False
        for region in grouped:
            if not grouped[region]:
                continue
            selected.append(grouped[region].pop(0))
            made_progress = True
            if len(selected) >= limit:
                break
        if not made_progress:
            break

    return selected

def decode_google_news_links(articles: list[Article]) -> list[Article]:
    pending = [article for article in articles if "news.google.com" in article.link]
    if not pending:
        return articles

    try:
        results = asyncio.run(
            gnews_decoder_async(
                [article.link for article in pending],
                timeout=15,
                concurrency=min(8, len(pending)),
            )
        )
    except Exception as exc:
        print(f"[WARN] Google News link decoding failed: {exc}")
        return articles

    if isinstance(results, dict):
        results = [results]

    decoded_links: dict[str, str] = {}
    for article, result in zip(pending, results):
        if result.get("success") and result.get("decoded_url"):
            decoded_links[article.link] = result["decoded_url"]
        else:
            print(f"[WARN] Could not decode link for {article.region}")

    return [
        replace(article, link=decoded_links.get(article.link, article.link))
        for article in articles
    ]


def extract_article_content(html_text: str, url: str) -> str:
    try:
        extracted = trafilatura.extract(
            html_text,
            url=url,
            include_comments=False,
            include_tables=False,
            favor_recall=True,
        )
    except Exception:
        extracted = None

    if extracted:
        return clean_text(extracted)

    try:
        metadata = trafilatura.extract_metadata(html_text)
        if metadata and metadata.description:
            return clean_text(metadata.description)
    except Exception:
        pass

    return ""


def fetch_article_content(article: Article) -> Article:
    try:
        response = requests.get(
            article.link,
            headers={
                "User-Agent": USER_AGENT,
                "Accept-Language": "en-US,en;q=0.9,zh-CN;q=0.8",
            },
            timeout=(5, 18),
            allow_redirects=True,
        )
        response.raise_for_status()
        content = extract_article_content(response.text[:800000], response.url)
        content = truncate_text(content, 1800)
        print(f"[CONTENT] {article.region}: {len(content)} chars")
        return replace(article, content=content)
    except Exception as exc:
        print(f"[WARN] Content fetch failed for {article.region}: {exc}")
        return article


def enrich_articles(articles: list[Article]) -> list[Article]:
    if not articles:
        return articles
    with ThreadPoolExecutor(max_workers=min(6, len(articles))) as executor:
        return list(executor.map(fetch_article_content, articles))


def parse_summary_response(content: str) -> dict[int, str]:
    text = content.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text)

    start = text.find("[")
    end = text.rfind("]")
    if start >= 0 and end > start:
        try:
            payload = json.loads(text[start : end + 1])
            summaries: dict[int, str] = {}
            for item in payload:
                summary = clean_text(item.get("summary"))
                if summary and chinese_char_count(summary) >= 4:
                    summaries[int(item["id"])] = summary
            if summaries:
                return summaries
        except Exception:
            pass

    summaries: dict[int, str] = {}
    pattern = re.compile(
        r'\{\s*"id"\s*:\s*(\d+)\s*,\s*"summary"\s*:\s*"((?:\\.|[^"\\])*)"\s*\}',
        re.DOTALL,
    )
    for item_id, encoded_summary in pattern.findall(text):
        try:
            summary = clean_text(json.loads(f'"{encoded_summary}"'))
        except Exception:
            summary = clean_text(encoded_summary)
        if summary and chinese_char_count(summary) >= 4:
            summaries[int(item_id)] = summary
    return summaries


def request_cloudflare_chat(messages: list[dict], max_tokens: int) -> str:
    if not CLOUDFLARE_ACCOUNT_ID or not CLOUDFLARE_API_TOKEN:
        raise RuntimeError("Cloudflare credentials are not configured")

    models = [
        item.strip()
        for item in os.getenv(
            "CLOUDFLARE_MODELS",
            "@cf/google/gemma-4-26b-a4b-it,"
            "@cf/aisingapore/gemma-sea-lion-v4-27b-it,"
            "@cf/qwen/qwen3.8-27b",
        ).split(",")
        if item.strip()
    ]
    url = (
        "https://api.cloudflare.com/client/v4/accounts/"
        f"{CLOUDFLARE_ACCOUNT_ID}/ai/v1/chat/completions"
    )
    last_error = None

    for model in models:
        try:
            print(f"[AI] Trying Cloudflare model: {model}")
            response = requests.post(
                url,
                headers={
                    "Authorization": f"Bearer {CLOUDFLARE_API_TOKEN}",
                    "Content-Type": "application/json",
                },
                json={
                    "model": model,
                    "messages": messages,
                    "temperature": 0.1,
                    "max_tokens": max_tokens,
                },
                timeout=(10, 180),
            )
            if response.status_code != 200:
                raise RuntimeError(
                    f"HTTP {response.status_code}: {response.text[:500]}"
                )
            data = response.json()
            if data.get("errors"):
                raise RuntimeError(str(data["errors"]))
            choices = data.get("choices") or []
            if not choices:
                raise RuntimeError("Cloudflare returned no choices")
            content = (choices[0].get("message") or {}).get("content")
            if isinstance(content, str) and content.strip():
                print(f"[AI] Cloudflare success with model: {model}")
                return content.strip()
            raise RuntimeError("Cloudflare returned empty content")
        except Exception as exc:
            last_error = exc
            print(f"[WARN] Cloudflare model failed ({model}): {exc}")

    raise RuntimeError("All Cloudflare models failed") from last_error


def request_openrouter_chat(messages: list[dict], max_tokens: int) -> str:
    if not OPENROUTER_API_KEY:
        raise RuntimeError("OpenRouter API key is not configured")

    last_error = None
    for model in SUMMARY_MODELS:
        try:
            print(f"[AI] Trying OpenRouter model: {model}")
            response = requests.post(
                OPENROUTER_URL,
                headers={
                    "Authorization": f"Bearer {OPENROUTER_API_KEY}",
                    "Content-Type": "application/json",
                    "HTTP-Referer": "https://github.com/forwardok43-jpg/sea-news-digest",
                    "X-Title": "SEA News Digest",
                },
                json={
                    "model": model,
                    "messages": messages,
                    "temperature": 0.1,
                    "max_tokens": max_tokens,
                },
                timeout=(10, 180),
            )
            if response.status_code != 200:
                raise RuntimeError(
                    f"HTTP {response.status_code}: {response.text[:500]}"
                )
            data = response.json()
            if data.get("error"):
                raise RuntimeError(str(data["error"]))
            choices = data.get("choices") or []
            if not choices:
                raise RuntimeError("OpenRouter returned no choices")
            message = choices[0].get("message", {}) or {}
            content = message.get("content") or message.get("reasoning")
            if isinstance(content, str) and content.strip():
                print(f"[AI] OpenRouter success with model: {model}")
                return content.strip()
            raise RuntimeError("OpenRouter returned empty content")
        except Exception as exc:
            last_error = exc
            print(f"[WARN] OpenRouter model failed ({model}): {exc}")

    raise RuntimeError("All OpenRouter models failed") from last_error


def call_ai_chat(messages: list[dict], max_tokens: int) -> str:
    errors: list[str] = []

    try:
        return request_cloudflare_chat(messages, max_tokens)
    except Exception as exc:
        errors.append(f"Cloudflare: {exc}")

    try:
        return request_openrouter_chat(messages, max_tokens)
    except Exception as exc:
        errors.append(f"OpenRouter: {exc}")

    raise RuntimeError("All AI providers failed: " + " | ".join(errors))


def request_article_summaries(articles: list[Article]) -> dict[int, str]:
    payload_items = []
    for local_id, article in enumerate(articles, start=1):
        payload_items.append(
            {
                "id": local_id,
                "country": article.source_country or article.region,
                "title": article.title,
                "source": article.publisher,
                "content": truncate_text(article.content or article.title, 700),
            }
        )

    system_prompt = (
        "你是专业的国际新闻中文编辑。请逐条阅读新闻，"
        "每条输出50到80个简体中文字的简要摘要。"
        "只写输入中出现的事实，包含主要主体、事件和影响。"
        "不评论、不推测、不添加外部信息。"
        "如果正文信息不足，只根据标题写最保守的事实摘要。"
        '只返回JSON数组，格式为：[{"id":1,"summary":"简要摘要"}]。'
    )
    user_prompt = json.dumps(payload_items, ensure_ascii=False)
    content = call_ai_chat(
        [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        max_tokens=2200,
    )
    summaries = parse_summary_response(content)
    if not summaries:
        raise RuntimeError(f"Could not parse article summaries: {content[:300]}")
    return summaries


def translate_title_to_chinese(title: str, source_lang: str) -> str:
    title = clean_text(title)
    if chinese_char_count(title) >= 6:
        return title

    language = (source_lang or "en").strip()
    language_pairs = [f"{language}|zh-CN"]
    if language != "en":
        language_pairs.append("en|zh-CN")

    for langpair in language_pairs:
        for attempt in range(3):
            try:
                response = requests.get(
                    "https://api.mymemory.translated.net/get",
                    params={"q": title[:450], "langpair": langpair},
                    timeout=20,
                )
                if response.status_code == 429:
                    time.sleep(2 * (attempt + 1))
                    continue
                response.raise_for_status()
                data = response.json()
                translated = clean_text(
                    (data.get("responseData") or {}).get("translatedText")
                )
                if chinese_char_count(translated) >= 4:
                    return translated
                break
            except Exception as exc:
                if attempt < 2:
                    time.sleep(2 * (attempt + 1))
                    continue
                print(f"[WARN] Title translation failed ({langpair}): {exc}")

    return ""

def summarize_articles(articles: list[Article]) -> list[Article]:
    summaries: dict[int, str] = {}

    translated_titles: dict[int, str] = {}
    with ThreadPoolExecutor(
        max_workers=max(1, min(TRANSLATE_WORKERS, len(articles)))
    ) as executor:
        futures = {
            executor.submit(
                translate_title_to_chinese,
                article.title,
                article.source_lang,
            ): index
            for index, article in enumerate(articles, start=1)
        }
        for future in as_completed(futures):
            index = futures[future]
            try:
                translated = future.result()
                if translated:
                    translated_titles[index] = translated
            except Exception as exc:
                print(f"[WARN] Title translation worker failed: {exc}")

    if ENABLE_AI_SUMMARY:
        batches = [
            (start, articles[start - 1 : start - 1 + SUMMARY_BATCH_SIZE])
            for start in range(1, len(articles) + 1, SUMMARY_BATCH_SIZE)
        ]
        with ThreadPoolExecutor(
            max_workers=max(1, min(SUMMARY_WORKERS, len(batches)))
        ) as executor:
            futures = {
                executor.submit(request_article_summaries, batch): start
                for start, batch in batches
            }
            for future in as_completed(futures):
                start = futures[future]
                try:
                    batch_summaries = future.result()
                    for local_id, summary in batch_summaries.items():
                        summaries[start + local_id - 1] = clean_text(summary)
                    print(
                        f"[SUMMARY] Batch starting at {start}: "
                        f"{len(batch_summaries)} summaries"
                    )
                except Exception as exc:
                    print(f"[WARN] Summary batch failed at {start}: {exc}")
    else:
        print("[INFO] AI content summaries are disabled")

    summarized: list[Article] = []
    for index, article in enumerate(articles, start=1):
        summary = clean_text(summaries.get(index, ""))
        if len(summary) > 120:
            summary = summary[:120].rstrip("。！？!?，,；;：:") + "。"
        summarized.append(
            replace(
                article,
                summary=summary,
                title_zh=translated_titles.get(index, ""),
            )
        )
    return summarized


def build_overview(articles: list[Article]) -> str:
    items = [
        article
        for article in articles
        if article.summary and chinese_char_count(article.summary) >= 10
    ]
    if len(items) < 5:
        print(f"[WARN] Not enough summaries for overview: {len(items)}")
        return ""

    lines = []
    for index, article in enumerate(items, start=1):
        lines.append(
            f"{index}. [{article.source_country or article.region}] "
            f"{article.summary}"
        )

    system_prompt = (
        "你是东盟与国际关系新闻主编。请根据提供的新闻摘要，"
        f"写一篇{OVERVIEW_MIN_CHARS}到{OVERVIEW_MAX_CHARS}个简体中文字的总体概述。"
        "优先总结东盟内部事务和成员国的重大政治、军事、经济新闻；"
        "其次总结东盟与主要大国的关系，以及世界经济与东盟的联系。"
        "删除重复信息，按主题自然组织，不写标题，不写编号，不评论，"
        "不得添加输入中没有的事实。"
    )
    overview = clean_text(
        call_ai_chat(
            [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": "\n".join(lines)},
            ],
            max_tokens=1800,
        )
    )
    overview = re.sub(r"^(总览|概述)[：:]\s*", "", overview)
    if len(overview) > OVERVIEW_MAX_CHARS:
        cutoff = overview[:OVERVIEW_MAX_CHARS]
        boundary = max(
            cutoff.rfind("。"),
            cutoff.rfind("！"),
            cutoff.rfind("？"),
        )
        if boundary >= OVERVIEW_MIN_CHARS:
            overview = cutoff[: boundary + 1]
        else:
            overview = cutoff.rstrip() + "。"

    print(
        f"[OVERVIEW] Generated {chinese_char_count(overview)} Chinese characters "
        f"from {len(items)} article summaries"
    )
    return overview

def build_fallback_overview(articles: list[Article]) -> str:
    if not articles:
        return ""

    entries = []
    for article in articles[:20]:
        country = article.source_country or article.region
        text = article.summary or article.title_zh or article.title
        entries.append(f"{country}：{text}")

    overview = "；".join(entries)
    if len(overview) > OVERVIEW_MAX_CHARS:
        cutoff = overview[:OVERVIEW_MAX_CHARS]
        boundary = max(cutoff.rfind("；"), cutoff.rfind("。"))
        overview = cutoff[: boundary + 1] if boundary > 0 else cutoff.rstrip() + "。"
    return overview


def configure_docx(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "Microsoft YaHei"
    )


def add_docx_run(
    paragraph,
    text: str,
    *,
    bold: bool = False,
    size: float = 10.5,
    color: tuple[int, int, int] | None = None,
) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "Microsoft YaHei"
    )
    if color:
        run.font.color.rgb = RGBColor(*color)


def create_overview_docx(
    overview: str,
    articles: list[Article],
    output_dir: Path,
) -> Path:
    now = datetime.now(TZ)
    document = Document()
    configure_docx(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_docx_run(title, "今日东盟新闻总览", bold=True, size=18)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_docx_run(
        meta,
        f"{now:%Y-%m-%d}｜覆盖 {len(articles)} 条新闻摘要",
        size=10,
        color=(90, 90, 90),
    )

    body = document.add_paragraph()
    body.paragraph_format.line_spacing = 1.5
    body.paragraph_format.space_before = Pt(10)
    add_docx_run(body, overview, size=11)

    path = output_dir / f"今日东盟新闻总览_{now:%Y-%m-%d}.docx"
    document.save(path)
    return path


def create_news_compilation_docx(
    articles: list[Article],
    output_dir: Path,
) -> Path:
    now = datetime.now(TZ)
    document = Document()
    configure_docx(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_docx_run(title, "东盟新闻摘要汇编", bold=True, size=18)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_docx_run(
        meta,
        f"{now:%Y-%m-%d %H:%M} 北京时间｜共 {len(articles)} 条",
        size=10,
        color=(90, 90, 90),
    )

    for index, article in enumerate(articles, start=1):
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(4)
        add_docx_run(
            heading,
            f"{index}. {article.region}｜{article.title_zh or article.title}",
            bold=True,
            size=12,
        )

        source_line = document.add_paragraph()
        source_line.paragraph_format.space_after = Pt(3)
        add_docx_run(
            source_line,
            f"出处国家/地区：{article.source_country or article.region}｜"
            f"来源：{source_label(article)}｜"
            f"时间：{article.published.astimezone(TZ):%Y-%m-%d %H:%M}",
            size=9,
            color=(100, 100, 100),
        )

        summary_line = document.add_paragraph()
        summary_line.paragraph_format.line_spacing = 1.35
        if article.summary:
            add_docx_run(summary_line, f"摘要：{article.summary}", size=10.5)
        else:
            add_docx_run(
                summary_line,
                "摘要：正文摘要暂未生成，请以中文标题为准。",
                size=10.5,
                color=(150, 70, 0),
            )

    path = output_dir / f"东盟新闻摘要汇编_{now:%Y-%m-%d}.docx"
    document.save(path)
    return path


def upload_wecom_file(session: requests.Session, path: Path) -> str:
    query = parse_qs(urlparse(WECOM_WEBHOOK_URL).query)
    key = (query.get("key") or [""])[0]
    if not key:
        raise RuntimeError("WECOM_WEBHOOK_URL does not contain key")

    upload_url = (
        "https://qyapi.weixin.qq.com/cgi-bin/webhook/upload_media"
        f"?key={quote(key)}&type=file"
    )
    last_error = None

    for attempt in range(3):
        try:
            with path.open("rb") as handle:
                response = session.post(
                    upload_url,
                    files={
                        "media": (
                            path.name,
                            handle,
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                    },
                    timeout=(10, 90),
                )
            response.raise_for_status()
            data = response.json()
            if data.get("errcode") != 0 or not data.get("media_id"):
                raise RuntimeError(f"WeCom upload error: {data}")
            return data["media_id"]
        except Exception as exc:
            last_error = exc
            if attempt < 2:
                time.sleep(2 * (attempt + 1))

    raise RuntimeError(f"Failed to upload Word document: {path.name}") from last_error


def send_wecom_file(session: requests.Session, media_id: str, filename: str) -> None:
    payload = {"msgtype": "file", "file": {"media_id": media_id}}
    response = session.post(WECOM_WEBHOOK_URL, json=payload, timeout=30)
    response.raise_for_status()
    data = response.json()
    if data.get("errcode") != 0:
        raise RuntimeError(f"WeCom file message error for {filename}: {data}")

def source_label(article: Article) -> str:
    if article.publisher:
        return article.publisher
    return urlparse(article.link).netloc.replace("www.", "")


def split_text(text: str, limit_bytes: int) -> list[str]:
    blocks = [block.strip() for block in text.split("\n\n") if block.strip()]
    messages: list[str] = []
    current = ""

    for block in blocks:
        candidate = f"{current}\n\n{block}" if current else block
        if current and len(candidate.encode("utf-8")) > limit_bytes:
            messages.append(current)
            current = block
        else:
            current = candidate

    if current:
        messages.append(current)
    return messages


def build_messages(articles: list[Article]) -> list[str]:
    now = datetime.now(TZ)
    parts = [
        "【东南亚新闻简报】",
        f"{now:%Y-%m-%d %H:%M} 北京时间",
        f"过去 {LOOKBACK_HOURS} 小时｜精选 {len(articles)} 条热点",
        "────────────────",
    ]

    article_blocks: list[str] = []
    for article in articles:
        local_time = article.published.astimezone(TZ).strftime("%m-%d %H:%M")
        title = article.title_zh or article.title
        block = f"【{article.region}】\n{title}"
        origin = article.source_country or article.region
        block += f"\n出处国家/地区：{origin}\n来源：{source_label(article)} · {local_time}"
        article_blocks.append(block)

    text = "\n\n".join(parts + article_blocks)
    text += (
        "\n\n────────────────\n"
        "说明：标题和摘要由免费翻译及摘要服务自动生成，重要信息请以原文为准。"
    )

    messages = split_text(text, 1900)
    if len(messages) > 1:
        messages = [messages[0]] + [
            f"【东南亚新闻简报（续）】\n\n{message}" for message in messages[1:]
        ]
    return messages

def send_wecom_text(session: requests.Session, content: str) -> None:
    payload = {"msgtype": "text", "text": {"content": content}}
    last_error = None

    for attempt in range(3):
        try:
            response = session.post(WECOM_WEBHOOK_URL, json=payload, timeout=25)
            response.raise_for_status()
            result = response.json()
            if result.get("errcode") != 0:
                raise RuntimeError(f"WeCom API error: {result}")
            return
        except Exception as exc:
            last_error = exc
            if attempt < 2:
                time.sleep(2 * (attempt + 1))

    raise RuntimeError("Failed to send WeCom message") from last_error


def main() -> None:
    if not DRY_RUN and not WECOM_WEBHOOK_URL:
        raise SystemExit("Missing WECOM_WEBHOOK_URL")

    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})
    articles, failures = collect_articles(session)

    if not articles:
        if failures:
            raise RuntimeError("All news sources failed:\n" + "\n".join(failures))
        print("No articles found in the requested time window")
        return

    digest_articles = select_digest_articles(articles, DIGEST_ITEMS)
    print(f"[DIGEST] Selected {len(digest_articles)} country-level articles")

    digest_articles = decode_google_news_links(digest_articles)
    if ENABLE_AI_SUMMARY:
        digest_articles = enrich_articles(digest_articles)
    digest_articles = summarize_articles(digest_articles)

    overview = build_overview(digest_articles) if ENABLE_AI_SUMMARY else ""
    if not overview:
        overview = build_fallback_overview(digest_articles)

    output_dir = Path("output")
    output_dir.mkdir(parents=True, exist_ok=True)
    overview_path = create_overview_docx(overview, digest_articles, output_dir)
    compilation_path = create_news_compilation_docx(digest_articles, output_dir)
    output_files = [overview_path, compilation_path]

    for path in output_files:
        print(f"[DOCX] {path} ({path.stat().st_size} bytes)")

    if DRY_RUN:
        print("[DRY RUN] Word documents generated but not sent")
        return

    system_session = requests.Session()
    system_session.headers.update({"User-Agent": USER_AGENT})

    for path in output_files:
        media_id = upload_wecom_file(system_session, path)
        send_wecom_file(system_session, media_id, path.name)
        print(f"[SEND FILE] {path.name}")
        time.sleep(2)

    print("Pushed 2 Word documents: 1 overview and 1 news-summary compilation")


if __name__ == "__main__":
    main()